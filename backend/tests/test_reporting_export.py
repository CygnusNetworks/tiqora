"""Pure tests of the report XLSX/CSV writers."""

from __future__ import annotations

import io
from datetime import datetime

import openpyxl

from tiqora.reporting.export import (
    XlsxColumn,
    XlsxSheet,
    report_sheet,
    write_report_csv,
    write_xlsx_workbook,
)
from tiqora.reporting.schemas import (
    FactsStatusOut,
    ReportResultOut,
    ReportTotalsOut,
    ResultColumnOut,
)


def _result(**over: object) -> ReportResultOut:
    base: dict[str, object] = {
        "report_key": "builtin.time_by_customer",
        "columns": [
            ResultColumnOut(
                key="customer",
                role="dimension",
                type="string",
                label_key="reports.dimensions.customer",
            ),
            ResultColumnOut(
                key="time_units",
                role="measure",
                type="units",
                label_key="reports.measures.time_units",
            ),
            ResultColumnOut(key="share", role="measure", type="percent", label_key="x"),
        ],
        "rows": [["ACME", 12.5, 0.625], ["__other__", 7.5, 0.375]],
        "labels": {"customer": {"ACME": "Acme GmbH"}},
        "totals": ReportTotalsOut(grand=[None, 20.0, 1.0]),
        "previous": None,
        "row_count": 2,
        "truncated": False,
        "time_window": None,
        "generated_at": datetime(2024, 6, 20, 12, 0),
        "cached": False,
        "facts": FactsStatusOut(
            ready=True,
            as_of=None,
            backfill_status="done",
            backfill_job_id=None,
            backfill_progress=None,
        ),
    }
    base.update(over)
    return ReportResultOut(**base)  # type: ignore[arg-type]


RESULT = _result()


def test_xlsx_structure() -> None:
    data = write_xlsx_workbook(
        [
            report_sheet(
                RESULT,
                title="Zeit nach Kunde: Juni/2024",
                column_labels={
                    "customer": "Kunde",
                    "time_units": "Zeiteinheiten",
                    "share": "Anteil",
                },
                value_labels={"customer": {"__other__": "Weitere"}},
                totals_label="Summe",
                info_lines=["Zeitraum: 01.06.2024 – 30.06.2024"],
            )
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Zeit nach Kunde Juni2024", "Info"]
    ws = wb["Zeit nach Kunde Juni2024"]
    assert [c.value for c in ws[1]] == ["Kunde", "Zeiteinheiten", "Anteil"] and ws["A1"].font.b
    assert ws["A1"].fill.fgColor.rgb.endswith("EDF0F7")
    assert ws.freeze_panes == "A2"
    assert [c.value for c in ws[2]] == ["Acme GmbH", 12.5, 0.625] and ws[
        "C2"
    ].number_format == "0.0%"
    assert ws["A3"].value == "Weitere" and ws["B2"].number_format == "0.00"
    assert [c.value for c in ws[4]] == ["Summe", 20.0, 1.0] and ws["A4"].font.b
    assert wb["Info"]["A1"].value == "Zeitraum: 01.06.2024 – 30.06.2024"


def test_none_measure_is_empty_and_label_precedence() -> None:
    res = _result(rows=[["ACME", None, 0.5], ["raw", 1, 0.5]], totals=None)
    sheet = report_sheet(
        res,
        title="T",
        column_labels={},
        value_labels={"customer": {"ACME": "Override"}},
        totals_label="S",
        info_lines=[],
    )
    ws = openpyxl.load_workbook(io.BytesIO(write_xlsx_workbook([sheet])))["T"]
    assert [c.value for c in ws[2]] == ["Override", None, 0.5]
    assert ws["A3"].value == "raw" and ws["A1"].value == "customer"
    assert ws.max_row == 3


def test_text_never_becomes_formula() -> None:
    res = _result(rows=[["=1+1", 1.0, 0.1]], totals=None)
    sheet = report_sheet(
        res, title="T", column_labels={}, value_labels={}, totals_label="S", info_lines=[]
    )
    ws = openpyxl.load_workbook(io.BytesIO(write_xlsx_workbook([sheet])))["T"]
    assert ws["A2"].value == "=1+1" and ws["A2"].data_type == "s"


def test_date_cells_are_real_dates() -> None:
    res = _result(
        columns=[
            ResultColumnOut(key="day", role="dimension", type="date", label_key="d"),
            ResultColumnOut(key="n", role="measure", type="count", label_key="n"),
        ],
        rows=[["2024-06-03", 4]],
        labels={},
        totals=None,
    )
    sheet = report_sheet(
        res, title="D", column_labels={}, value_labels={}, totals_label="S", info_lines=[]
    )
    ws = openpyxl.load_workbook(io.BytesIO(write_xlsx_workbook([sheet])))["D"]
    assert ws["A2"].value == datetime(2024, 6, 3)
    assert ws["A2"].number_format == "yyyy-mm-dd" and ws["B2"].number_format == "0"


def test_sheet_names_are_unique_and_short() -> None:
    cols = [XlsxColumn("a", "string")]
    name = "x" * 40
    data = write_xlsx_workbook(
        [
            XlsxSheet(name=name, columns=cols, rows=[["v"]]),
            XlsxSheet(name=name, columns=cols, rows=[["w"]]),
        ]
    )
    assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames == ["x" * 31, "x" * 27 + " (2)"]


def test_widths_and_streamed_rows_beyond_scan() -> None:
    cols = [XlsxColumn("a", "string"), XlsxColumn("b", "count", width=33.0)]
    rows = ([f"row-{i}" + "y" * 80 if i == 0 else f"r{i}", i] for i in range(700))
    ws = openpyxl.load_workbook(
        io.BytesIO(write_xlsx_workbook([XlsxSheet(name="W", columns=cols, rows=rows)]))
    )["W"]
    assert ws.max_row == 701 and ws["B701"].value == 699
    assert ws.column_dimensions["A"].width == 60 and ws.column_dimensions["B"].width == 33


def test_csv_has_bom_semicolons_and_labels() -> None:
    data = write_report_csv(
        RESULT, column_labels={"customer": "Kunde"}, value_labels={}, totals_label="Summe"
    )
    assert data.startswith(b"\xef\xbb\xbf")
    lines = data.decode("utf-8-sig").splitlines()
    assert lines[0] == "Kunde;time_units;share" and lines[1] == "Acme GmbH;12.5;0.625"
    assert lines[2] == "Weitere;7.5;0.375" and lines[-1] == "Summe;20.0;1.0"


def test_csv_none_empty_and_formula_text_neutralised() -> None:
    res = _result(rows=[["=cmd", None, 0.5]], totals=None)
    lines = (
        write_report_csv(res, column_labels={}, value_labels={}, totals_label="S")
        .decode("utf-8-sig")
        .splitlines()
    )
    assert lines[1] == "'=cmd;;0.5" and len(lines) == 2


def _ws(res: ReportResultOut, title: str = "T"):  # type: ignore[no-untyped-def]
    sheet = report_sheet(
        res, title=title, column_labels={}, value_labels={}, totals_label="Summe", info_lines=[]
    )
    return openpyxl.load_workbook(io.BytesIO(write_xlsx_workbook([sheet])))[title]


def _date_result(rows: list[list[object]], totals: object = None) -> ReportResultOut:
    return _result(
        columns=[
            ResultColumnOut(key="day", role="dimension", type="date", label_key="d"),
            ResultColumnOut(key="n", role="measure", type="count", label_key="n"),
        ],
        rows=rows,
        labels={},
        totals=totals,
    )


def test_tz_aware_dates_become_naive_utc() -> None:
    ws = _ws(_date_result([["2024-06-03T10:00:00Z", 1], ["2024-06-03T12:30:00+02:00", 2]]), "D")
    assert ws["A2"].value == datetime(2024, 6, 3, 10, 0)
    assert ws["A3"].value == datetime(2024, 6, 3, 10, 30)


def test_datetime_with_time_part_keeps_minutes_format() -> None:
    ws = _ws(_date_result([["2024-06-03 10:15:00", 1], ["2024-06-04", 2]]), "D")
    assert ws["A2"].number_format == "yyyy-mm-dd hh:mm" and ws["A2"].value == datetime(
        2024, 6, 3, 10, 15
    )
    assert ws["A3"].number_format == "yyyy-mm-dd"


def test_totals_label_goes_into_first_dimension_column() -> None:
    res = _result(
        columns=[
            ResultColumnOut(key="n", role="measure", type="count", label_key="n"),
            ResultColumnOut(key="customer", role="dimension", type="string", label_key="c"),
        ],
        rows=[[1, "ACME"]],
        totals=ReportTotalsOut(grand=[None, None]),
    )
    ws = _ws(res)
    assert [c.value for c in ws[3]] == [None, "Summe"]


def test_totals_label_for_measure_only_report_adds_label_column() -> None:
    res = _result(
        columns=[ResultColumnOut(key="n", role="measure", type="count", label_key="n")],
        rows=[[3]],
        totals=ReportTotalsOut(grand=[3]),
    )
    ws = _ws(res)
    assert [c.value for c in ws[3]] == ["Summe", 3] and ws["B2"].value == 3
    lines = (
        write_report_csv(res, column_labels={}, value_labels={}, totals_label="Summe")
        .decode("utf-8-sig")
        .splitlines()
    )
    assert lines == [";n", ";3", "Summe;3"]


def test_other_in_date_dimension_is_weitere() -> None:
    ws = _ws(_date_result([["2024-06-03", 1], ["__other__", 2]]), "D")
    assert ws["A3"].value == "Weitere"
    csv = write_report_csv(
        _date_result([["__other__", 2]]), column_labels={}, value_labels={}, totals_label="S"
    )
    assert csv.decode("utf-8-sig").splitlines()[1] == "Weitere;2"


def test_nan_and_inf_are_empty() -> None:
    res = _result(rows=[["ACME", float("nan"), float("inf")]], totals=None)
    ws = _ws(res)
    assert [c.value for c in ws[2]] == ["Acme GmbH", None, None]
    lines = (
        write_report_csv(res, column_labels={}, value_labels={}, totals_label="S")
        .decode("utf-8-sig")
        .splitlines()
    )
    assert lines[1] == "Acme GmbH;;"


def test_sheet_names_trim_and_avoid_history() -> None:
    cols = [XlsxColumn("a", "string")]
    names = ["a" * 30 + " b", "History", "x" * 29 + "'y"]
    data = write_xlsx_workbook([XlsxSheet(name=n, columns=cols, rows=[]) for n in names])
    out = openpyxl.load_workbook(io.BytesIO(data)).sheetnames
    assert out[0] == "a" * 30 and out[1] == "History_" and out[2] == "x" * 29
    assert all(not n.endswith((" ", "'")) for n in out)


def test_csv_guard_prefixes_numeric_looking_strings_on_purpose() -> None:
    res = _result(rows=[["-5", 1.0, 0.1], ["+49 123", 1.0, 0.1]], totals=None)
    lines = (
        write_report_csv(res, column_labels={}, value_labels={}, totals_label="S")
        .decode("utf-8-sig")
        .splitlines()
    )
    assert lines[1].startswith("'-5;") and lines[2].startswith("'+49 123;")
