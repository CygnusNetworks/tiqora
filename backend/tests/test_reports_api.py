"""HTTP tests for ``/api/v1/reports*`` and ``/api/v1/reporting/*`` plus the run/export jobs.

Both dialects. Seed ids 47810-47869 (names contain ``rep47``): agent 47810 (``ro`` on group
47811), readable queue 47811, foreign queue 47812; tickets 47820/47821 (readable queue, June
2024, e-mail) and 47822 (foreign queue); bookings 47830-47832 for customer ``rep47-cust16``.
The result cache is an in-memory backend shared by the API (``get_report_cache`` override)
and the jobs (``tiqora.reporting.jobs.cache_backend_from_settings``). Every row this module
writes is deleted again (incl. its jobs, the report state row and ``reporting.*`` settings).
"""

from __future__ import annotations

import asyncio
import io
import time
from collections.abc import AsyncIterator, Generator
from datetime import datetime
from decimal import Decimal
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

import tiqora.api.v1.reports as reports_api
import tiqora.reporting.jobs as report_jobs
import tiqora.reporting.run_service as run_service
from tiqora.api.app import create_app
from tiqora.api.deps import get_current_user, get_db
from tiqora.api.v1.reports import get_report_cache
from tiqora.config import Settings
from tiqora.db.tiqora.base import TiqoraBase
from tiqora.domain.auth import AuthenticatedUser
from tiqora.extensions import get_registry
from tiqora.jobs.runner import JobRunner
from tiqora.reporting.cache import GENERATION_KEY
from tiqora.reporting.engine import QueryBudget
from tiqora.reporting.facts_store import rebuild_ticket_facts
from tiqora.reporting.settings import REPORTING_SECTION
from tiqora.reporting.sql import ReportQueryTimeout
from tiqora.reporting.state import load_state
from tiqora.settings_sections.store import save_section_values

pytestmark = pytest.mark.db

LO, HI = 47810, 47869
AGENT, QUEUE, FOREIGN = 47810, 47811, 47812
T1, T2, T_FOREIGN = 47820, 47821, 47822
TICKETS = (T1, T2, T_FOREIGN)
CUSTOMER = "rep47-cust16"
CUSTOMER_USER = "rep47.cust16@example.org"
BUILT = datetime(2024, 6, 20, 11, 0)
JUNE = {"kind": "absolute", "start": "2024-06-01", "end": "2024-06-30"}
TIMEOUT_TEXT = (
    "Der Bericht hat zu lange gebraucht. "
    "Bitte den Zeitraum verkleinern oder ihn als Export starten."
)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
USER = AuthenticatedUser(
    id=AGENT, login="rep47.api", first_name="Rhea", last_name="Reports", auth_method="session"
)
#: Znuny seed admin (group "admin" rw)
ADMIN = AuthenticatedUser(
    id=1, login="root@localhost", first_name="Admin", last_name="OTRS", auth_method="session"
)
_BAND = {"lo": LO, "hi": HI}
_JOBS = "kind IN ('core.report_run', 'core.report_export')"
_RESET = (
    f"DELETE FROM tiqora_job_file WHERE job_id IN (SELECT id FROM tiqora_job WHERE {_JOBS})",
    f"DELETE FROM tiqora_job WHERE {_JOBS}",
    "DELETE FROM tiqora_report_state",
    "DELETE FROM tiqora_settings WHERE {key} LIKE 'reporting.%'",
)
_CLEANUP = (
    *_RESET,
    "DELETE FROM tiqora_report_ticket_event WHERE ticket_id BETWEEN :lo AND :hi",
    "DELETE FROM tiqora_report_ticket_fact WHERE ticket_id BETWEEN :lo AND :hi",
    "DELETE FROM tiqora_report_dirty_ticket WHERE ticket_id BETWEEN :lo AND :hi",
    "DELETE FROM time_accounting WHERE ticket_id BETWEEN :lo AND :hi",
    "DELETE FROM ticket_history WHERE ticket_id BETWEEN :lo AND :hi",
    "DELETE FROM article WHERE ticket_id BETWEEN :lo AND :hi",
    "DELETE FROM ticket WHERE id BETWEEN :lo AND :hi",
    "DELETE FROM group_user WHERE user_id BETWEEN :lo AND :hi",
    "DELETE FROM queue WHERE id BETWEEN :lo AND :hi",
    "DELETE FROM permission_groups WHERE id BETWEEN :lo AND :hi",
    "DELETE FROM users WHERE id BETWEEN :lo AND :hi",
)
_URLS: set[str] = set()
_SEEDED: set[str] = set()


def _to_async_url(sync_url: str) -> str:
    if sync_url.startswith("postgresql+psycopg2://"):
        return sync_url.replace("postgresql+psycopg2://", "postgresql+asyncpg://", 1)
    if sync_url.startswith("postgresql://"):
        return sync_url.replace("postgresql://", "postgresql+asyncpg://", 1)
    if sync_url.startswith("mysql+pymysql://"):
        return sync_url.replace("mysql+pymysql://", "mysql+aiomysql://", 1)
    return sync_url


def _run_sql(sync_url: str, statements: tuple[str, ...]) -> None:
    engine = create_engine(sync_url)
    TiqoraBase.metadata.create_all(engine)
    key = "`key`" if engine.dialect.name == "mysql" else '"key"'
    with engine.begin() as conn:
        for stmt in statements:
            conn.execute(text(stmt.replace("{key}", key)), _BAND)
    engine.dispose()


@pytest.fixture(scope="module", autouse=True)
def _delete_seeded_rows() -> Generator[None, None, None]:
    yield
    for url in _URLS:
        _run_sql(url, _CLEANUP)
    _URLS.clear()
    _SEEDED.clear()


def _id(conn: Any, sql: str, **params: Any) -> int:
    return int(conn.execute(text(sql), params).scalar_one())


def _grant(conn: Any) -> None:
    conn.execute(
        text(
            "INSERT INTO group_user (user_id, group_id, permission_key,"
            " create_time, create_by, change_time, change_by)"
            " VALUES (:u, :g, 'ro', :t, 1, :t, 1)"
        ),
        {"u": AGENT, "g": QUEUE, "t": BUILT},
    )


def _seed(sync_url: str) -> None:
    t0 = datetime(2024, 5, 1)
    engine = create_engine(sync_url)
    with engine.begin() as conn:
        h_new = _id(conn, "SELECT id FROM ticket_history_type WHERE name = 'NewTicket'")
        s_customer = _id(conn, "SELECT id FROM article_sender_type WHERE name = 'customer'")
        chan = _id(conn, "SELECT id FROM communication_channel WHERE name = 'Email'")
        st_new = _id(conn, "SELECT id FROM ticket_state WHERE name = 'new'")
        conn.execute(
            text(
                "INSERT INTO users (id, login, pw, first_name, last_name, valid_id,"
                " create_time, create_by, change_time, change_by)"
                " VALUES (:id, 'rep47.api', 'x', 'Rhea', 'Reports', 1, :t, 1, :t, 1)"
            ),
            {"id": AGENT, "t": t0},
        )
        for gid in (QUEUE, FOREIGN):
            conn.execute(
                text(
                    "INSERT INTO permission_groups (id, name, valid_id,"
                    " create_time, create_by, change_time, change_by)"
                    " VALUES (:id, :n, 1, :t, 1, :t, 1)"
                ),
                {"id": gid, "n": f"rep47-group-{gid}", "t": t0},
            )
            conn.execute(
                text(
                    "INSERT INTO queue (id, name, group_id, system_address_id, salutation_id,"
                    " signature_id, follow_up_id, follow_up_lock, valid_id,"
                    " create_time, create_by, change_time, change_by)"
                    " VALUES (:id, :n, :id, 1, 1, 1, 1, 0, 1, :t, 1, :t, 1)"
                ),
                {"id": gid, "n": f"rep47-queue-{gid}", "t": t0},
            )
        _grant(conn)
        for n, (tid, queue, at) in enumerate(
            (
                (T1, QUEUE, datetime(2024, 6, 3, 8)),
                (T2, QUEUE, datetime(2024, 6, 4, 9)),
                (T_FOREIGN, FOREIGN, datetime(2024, 6, 5, 10)),
            )
        ):
            conn.execute(
                text(
                    "INSERT INTO ticket (id, tn, title, queue_id, ticket_lock_id, type_id,"
                    " user_id, responsible_user_id, ticket_priority_id, ticket_state_id,"
                    " customer_id, customer_user_id, timeout, until_time, escalation_time,"
                    " escalation_update_time, escalation_response_time,"
                    " escalation_solution_time, archive_flag,"
                    " create_time, create_by, change_time, change_by)"
                    " VALUES (:id, :tn, :title, :q, 1, 1, :u, 1, 3, :s, :c, :cu,"
                    " 0, 0, 0, 0, 0, 0, 0, :t, 1, :t, 1)"
                ),
                {"id": tid, "tn": f"20240600{tid}", "title": f"rep47 ticket {tid}", "q": queue,
                 "u": AGENT, "s": st_new, "c": CUSTOMER, "cu": CUSTOMER_USER, "t": at},
            )  # fmt: skip
            conn.execute(
                text(
                    "INSERT INTO ticket_history (id, name, history_type_id, ticket_id,"
                    " article_id, type_id, queue_id, owner_id, priority_id, state_id,"
                    " create_time, create_by, change_time, change_by)"
                    " VALUES (:id, :n, :ht, :tid, NULL, 1, :q, :o, 3, :s, :t, :o, :t, :o)"
                ),
                {"id": LO + 30 + n, "n": f"%%20240600{tid}%%rep47-queue-{queue}%%3 normal%%new",
                 "ht": h_new, "tid": tid, "q": queue, "o": AGENT, "s": st_new, "t": at},
            )  # fmt: skip
            conn.execute(
                text(
                    "INSERT INTO article (id, ticket_id, article_sender_type_id,"
                    " communication_channel_id, is_visible_for_customer,"
                    " create_time, create_by, change_time, change_by)"
                    " VALUES (:id, :tid, :s, :c, 1, :t, 1, :t, 1)"
                ),
                {"id": LO + 30 + n, "tid": tid, "s": s_customer, "c": chan, "t": at},
            )
        for taid, tid, units, at in (
            (47830, T1, "1.50", datetime(2024, 6, 5, 10)),
            (47831, T2, "2.00", datetime(2024, 6, 6, 10)),
            (47832, T_FOREIGN, "3.00", datetime(2024, 6, 7, 10)),
        ):
            conn.execute(
                text(
                    "INSERT INTO time_accounting (id, ticket_id, article_id, time_unit,"
                    " create_time, create_by, change_time, change_by)"
                    " VALUES (:id, :tid, NULL, :u, :t, :by, :t, :by)"
                ),
                {"id": taid, "tid": tid, "u": Decimal(units), "t": at, "by": AGENT},
            )
    engine.dispose()


class _MemoryCache:
    """In-memory stand-in for the Redis result cache (``get`` / ``set(ex=)``)."""

    def __init__(self) -> None:
        self.store: dict[str, str] = {}

    async def get(self, key: str) -> str | None:
        return self.store.get(key)

    async def set(self, key: str, value: str, ex: int) -> None:
        self.store[key] = value


class Env:
    def __init__(
        self, sync_url: str, factory: async_sessionmaker[AsyncSession], client: AsyncClient
    ) -> None:
        self.sync_url = sync_url
        self.factory = factory
        self.client = client
        #: the user every request is made as (``get_current_user`` override)
        self.user = USER
        #: the in-memory result cache shared by the API and the jobs
        self.cache = _MemoryCache()

    @property
    def dialect(self) -> str:
        return "postgresql" if self.sync_url.startswith("postgresql") else "mysql"

    async def run_jobs(self) -> None:
        runner = JobRunner(
            session_factory=self.factory,
            settings=Settings(environment="test"),
            registry=get_registry(),
        )
        await runner.run_once()
        await runner.wait_idle()

    async def job(self, job_id: int) -> dict[str, Any]:
        resp = await self.client.get(f"/api/v1/jobs/{job_id}")
        assert resp.status_code == 200, resp.text
        return resp.json()  # type: ignore[no-any-return]

    async def ensure_state(self) -> None:
        """Commit the state row, as the worker's first tick does in production."""
        async with self.factory() as s:
            await load_state(s)
            await s.commit()

    async def facts_ready(self) -> None:
        async with self.factory() as s:
            (await load_state(s)).backfill_status = "done"
            await s.commit()

    async def set_access(self, access: str) -> None:
        async with self.factory() as s:
            await save_section_values(
                s,
                REPORTING_SECTION,
                {"access": access},
                secret_key=Settings(environment="test").secret_key,
            )
            await s.commit()


@pytest.fixture(params=["mariadb_znuny_url", "postgres_znuny_url"])
async def env(
    request: pytest.FixtureRequest, monkeypatch: pytest.MonkeyPatch
) -> AsyncIterator[Env]:
    sync_url: str = request.getfixturevalue(request.param)
    _URLS.add(sync_url)
    engine = create_async_engine(_to_async_url(sync_url))
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    if sync_url not in _SEEDED:
        _run_sql(sync_url, _CLEANUP)
        _seed(sync_url)
        async with factory() as s:
            await rebuild_ticket_facts(s, TICKETS, now=BUILT)
            await s.commit()
        _SEEDED.add(sync_url)
    _run_sql(sync_url, _RESET)

    cache = _MemoryCache()
    monkeypatch.setattr(report_jobs, "cache_backend_from_settings", lambda _settings: cache)

    async def _db() -> Any:
        async with factory() as session:
            yield session

    app = create_app(Settings(environment="test"))
    holder: list[Env] = []
    app.dependency_overrides[get_current_user] = lambda: holder[0].user
    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_report_cache] = lambda: cache
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
            holder.append(Env(sync_url, factory, c))
            holder[0].cache = cache
            yield holder[0]
    finally:
        await engine.dispose()
        _run_sql(sync_url, _RESET)


async def _run(env: Env, body: dict[str, Any]) -> Any:
    return await env.client.post("/api/v1/reports/run", json=body)


# ---------------------------------------------------------------- catalogue, sources


async def test_catalogue_lists_builtins_and_facts_status(env: Env) -> None:
    resp = await env.client.get("/api/v1/reports")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert len(body["reports"]) == 25
    first = body["reports"][0]
    assert first["key"] == "builtin.ticket_kpis"
    assert first["category"] == "overview" and first["overview"] is True
    assert first["kind"] == "builtin" and first["read_only"] is True
    assert first["id"] is None and first["group_ids"] == [] and first["owner"] is None
    assert {r["kind"] for r in body["reports"]} == {"builtin"}  # no saved rows seeded here
    assert first["title_key"] == "reports.builtin.ticket_kpis.title"
    assert body["facts"]["backfill_status"] in ("pending", "running", "done")
    assert body["facts"]["ready"] is False  # state row reset per test


async def test_get_report_returns_the_catalogue_definition(env: Env) -> None:
    resp = await env.client.get("/api/v1/reports/builtin.volume_by_queue")
    assert resp.status_code == 200
    definition = resp.json()["definition"]
    assert definition["source"] == "tickets"
    # parameter controls come from the original definition: default/required untouched
    assert definition["parameters"][0]["key"] == "queues"
    assert definition["parameters"][0]["default"] is None
    assert definition["dimensions"][0]["choices"] == ["type", "service", "priority"]


async def test_get_unknown_report_404(env: Env) -> None:
    resp = await env.client.get("/api/v1/reports/builtin.nope")
    assert resp.status_code == 404
    assert "builtin.nope" in resp.json()["detail"]["message"]
    run = await _run(env, {"report_key": "builtin.nope"})
    assert run.status_code == 404


async def test_sources_endpoint(env: Env) -> None:
    resp = await env.client.get("/api/v1/reporting/sources")
    assert resp.status_code == 200
    sources = {s["key"]: s for s in resp.json()}
    assert set(sources) == {
        "tickets",
        "ticket_events",
        "daily_snapshot",
        "accounted_time",
        "open_tickets",
    }
    tickets = sources["tickets"]
    queue = next(d for d in tickets["dimensions"] if d["key"] == "queue")
    assert queue["options_source"] == "queue" and queue["filter_type"] == "id"
    assert queue["labels"] == "queue" and queue["time_field"] is False
    created = next(d for d in tickets["dimensions"] if d["key"] == tickets["default_time_field"])
    assert created["time_field"] is True
    assert any(m["key"] == "tickets" and m["agg"] == "count" for m in tickets["measures"])
    for src in sources.values():
        for f in src["filter_fields"]:
            assert f["kind"] == f["filter_type"] and f["time_field"] is False


async def test_status_endpoint(env: Env) -> None:
    resp = await env.client.get("/api/v1/reporting/status")
    assert resp.status_code == 200
    assert resp.json()["ready"] is False and resp.json()["backfill_status"] == "pending"
    await env.facts_ready()
    assert (await env.client.get("/api/v1/reporting/status")).json()["ready"] is True


# ---------------------------------------------------------------- inline runs


async def test_run_builtin_inline(env: Env) -> None:
    resp = await _run(env, {"report_key": "builtin.channel_mix", "time_range": JUNE})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["job"] is None
    result = body["result"]
    assert result["report_key"] == "builtin.channel_mix"
    assert result["columns"][0]["key"] == "channel"
    assert result["rows"] == [["email", 2]]  # the foreign queue's ticket is not counted
    assert result["cached"] is False
    window = result["time_window"]
    assert window["start"].endswith(("Z", "+00:00"))
    assert window["start_date"] == "2024-06-01" and window["end_date"] == "2024-06-30"


async def test_run_inline_definition(env: Env) -> None:
    resp = await _run(
        env,
        {"definition": {"source": "tickets", "measures": [{"key": "tickets"}]}, "time_range": JUNE},
    )
    assert resp.status_code == 200, resp.text
    result = resp.json()["result"]
    assert result["report_key"] is None
    assert result["rows"] == [[2]]


async def test_run_time_zone_order(env: Env) -> None:
    body = {"definition": {"source": "tickets", "measures": [{"key": "tickets"}],
                           "time_zone": "America/New_York"}, "time_range": JUNE}  # fmt: skip
    from_definition = (await _run(env, body)).json()["result"]["time_window"]
    assert from_definition["time_zone"] == "America/New_York"
    from_request = (await _run(env, {**body, "time_zone": "Europe/Berlin"})).json()["result"]
    assert from_request["time_window"]["time_zone"] == "Europe/Berlin"


async def test_run_requires_exactly_one_of_key_or_definition(env: Env) -> None:
    definition = {"source": "tickets", "measures": [{"key": "tickets"}]}
    both = await _run(env, {"report_key": "builtin.channel_mix", "definition": definition})
    neither = await _run(env, {"time_range": JUNE})
    assert both.status_code == 422 and neither.status_code == 422
    assert both.json()["detail"]["path"] == "report_key"


async def test_invalid_definition_422_with_path(env: Env) -> None:
    body = {"definition": {"source": "tickets", "measures": [{"key": "nope"}]}}
    resp = await _run(env, body)
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert set(detail) == {"message", "path"}
    assert detail["path"] == "measures[0].key" and "nope" in detail["message"]
    export = await env.client.post("/api/v1/reports/export", json=body)
    assert export.status_code == 422 and export.json()["detail"]["path"] == "measures[0].key"


async def test_invalid_parameter_is_422(env: Env) -> None:
    resp = await _run(env, {"report_key": "builtin.channel_mix", "parameters": {"queues": ["x"]}})
    assert resp.status_code == 422
    assert resp.json()["detail"]["path"] == "parameters.queues"


async def test_internal_run_fields_are_rejected(env: Env) -> None:
    window = {"start_utc": "2020-01-01T00:00:00", "end_utc": "2030-01-01T00:00:00"}
    for field, value in (
        ("window_override", window),
        ("extra_filter", {"where": {"field": "queue", "op": "eq", "value": FOREIGN}}),
        ("last_snapshot_date", "2024-06-01"),
    ):
        body = {"report_key": "builtin.channel_mix", field: value}
        run = await _run(env, body)
        assert run.status_code == 422, (field, run.text)
        assert any(field in err["loc"] for err in run.json()["detail"])
        export = await env.client.post("/api/v1/reports/export", json=body)
        assert export.status_code == 422, field


async def test_dimension_choices_keys_are_indices(env: Env) -> None:
    ok = await _run(
        env,
        {"report_key": "builtin.volume_by_queue", "time_range": JUNE,
         "dimension_choices": {"0": "type"}},
    )  # fmt: skip
    assert ok.status_code == 200, ok.text
    assert ok.json()["result"]["columns"][0]["key"] == "type"
    bad = await _run(
        env, {"report_key": "builtin.volume_by_queue", "dimension_choices": {"x": "type"}}
    )
    assert bad.status_code == 422
    assert bad.json()["detail"]["path"] == "dimension_choices.x"
    unknown = await _run(
        env, {"report_key": "builtin.volume_by_queue", "dimension_choices": {"0": "nope"}}
    )
    assert unknown.status_code == 422


async def test_inline_budget_is_shared_by_estimate_and_run(
    env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    seen: list[QueryBudget | None] = []
    real_estimate, real_run = run_service.estimate_rows, run_service.execute_and_store

    async def estimate(*args: Any, **kwargs: Any) -> int:
        seen.append(kwargs.get("budget"))
        return await real_estimate(*args, **kwargs)

    async def run(*args: Any, **kwargs: Any) -> Any:
        seen.append(kwargs.get("budget"))
        return await real_run(*args, **kwargs)

    monkeypatch.setattr(run_service, "estimate_rows", estimate)
    monkeypatch.setattr(run_service, "execute_and_store", run)
    resp = await _run(env, {"report_key": "builtin.channel_mix", "time_range": JUNE})
    assert resp.status_code == 200
    assert len(seen) == 2 and seen[0] is not None and seen[0] is seen[1]
    budget = seen[0]
    assert budget.deadline is not None
    assert 0 < budget.deadline - budget.clock() <= 75


async def test_timeout_is_504(env: Env, monkeypatch: pytest.MonkeyPatch) -> None:
    async def slow(*args: Any, **kwargs: Any) -> Any:
        raise ReportQueryTimeout("budget used up")

    monkeypatch.setattr(run_service, "execute_and_store", slow)
    resp = await _run(env, {"report_key": "builtin.channel_mix", "time_range": JUNE})
    assert resp.status_code == 504
    assert resp.json() == {"detail": {"message": TIMEOUT_TEXT}}


# ---------------------------------------------------------------- background runs


async def test_force_job_returns_202_and_job_result_is_report(env: Env) -> None:
    await env.facts_ready()  # results are cached only when the report data is ready
    body = {"report_key": "builtin.channel_mix", "time_range": JUNE}
    resp = await _run(env, {**body, "force_job": True})
    assert resp.status_code == 202, resp.text
    accepted = resp.json()
    assert accepted["result"] is None
    job_id = accepted["job"]["job_id"]
    assert accepted["job"]["status"] == "queued"
    await env.run_jobs()
    job = await env.job(job_id)
    assert job["status"] == "succeeded", job
    assert job["kind"] == "core.report_run"
    result = job["result"]
    assert [c["key"] for c in result["columns"]] == ["channel", "tickets"]
    assert result["rows"] == [["email", 2]] and result["cached"] is False
    # the job wrote the cache: the same request is now answered inline from it
    again = await _run(env, body)
    assert again.status_code == 200
    assert again.json()["result"]["cached"] is True
    assert again.json()["result"]["rows"] == [["email", 2]]
    fresh = await _run(env, {**body, "no_cache": True})
    assert fresh.status_code == 200 and fresh.json()["result"]["cached"] is False


async def test_purge_during_an_inline_run_leaves_no_findable_result(
    env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    """T13 M2: the request reads the cache generation first; a purge (GDPR erasure,
    rebuild) while it computes makes its stored result unreachable."""
    await env.facts_ready()
    body = {"report_key": "builtin.channel_mix", "time_range": JUNE}
    real_estimate = run_service.estimate_rows

    async def estimate_then_purge(*args: Any, **kwargs: Any) -> int:
        rows = await real_estimate(*args, **kwargs)
        env.cache.store[GENERATION_KEY] = "1"  # what purge_report_cache's INCR does
        return rows

    monkeypatch.setattr(run_service, "estimate_rows", estimate_then_purge)
    first = await _run(env, body)
    assert first.status_code == 200 and first.json()["result"]["cached"] is False
    assert any(k.endswith(":g0") for k in env.cache.store)  # stored under the old generation
    monkeypatch.setattr(run_service, "estimate_rows", real_estimate)
    again = await _run(env, body)
    assert again.json()["result"]["cached"] is False  # not found: computed afresh ...
    third = await _run(env, body)
    assert third.json()["result"]["cached"] is True  # ... and cached under generation 1
    assert any(k.endswith(":g1") for k in env.cache.store)


async def test_threshold_triggers_job(env: Env, monkeypatch: pytest.MonkeyPatch) -> None:
    async def huge(*args: Any, **kwargs: Any) -> int:
        return 10**7

    monkeypatch.setattr(run_service, "estimate_rows", huge)
    resp = await _run(env, {"report_key": "builtin.channel_mix", "time_range": JUNE})
    assert resp.status_code == 202
    assert resp.json()["job"]["job_id"] > 0


async def test_long_event_window_runs_as_job(env: Env) -> None:
    definition = {"source": "ticket_events", "measures": [{"key": "events"}]}
    long_range = {"kind": "absolute", "start": "2023-01-01", "end": "2024-01-02"}
    resp = await _run(env, {"definition": definition, "time_range": long_range})
    assert resp.status_code == 202, resp.text
    short_range = {"kind": "absolute", "start": "2023-01-01", "end": "2024-01-01"}  # 366 days
    short = await _run(env, {"definition": definition, "time_range": short_range})
    assert short.status_code == 200, short.text


async def test_enqueue_runs_without_the_estimate_statement_timeout(
    env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    """PostgreSQL keeps the estimate's ``SET LOCAL statement_timeout`` until the transaction
    ends; the job row must be written in a fresh transaction without it."""
    real_estimate, real_enqueue = run_service.estimate_rows, run_service.enqueue_job
    timeouts: list[Any] = []

    async def estimate(*args: Any, **kwargs: Any) -> int:
        await real_estimate(*args, **kwargs)  # sets the statement timeout
        return 10**7

    async def enqueue(session: AsyncSession, **kwargs: Any) -> Any:
        sql = (
            "SHOW statement_timeout"
            if env.dialect == "postgresql"
            else "SELECT @@SESSION.max_statement_time"
        )
        timeouts.append((await session.execute(text(sql))).scalar_one())
        return await real_enqueue(session, **kwargs)

    monkeypatch.setattr(run_service, "estimate_rows", estimate)
    monkeypatch.setattr(run_service, "enqueue_job", enqueue)
    resp = await _run(env, {"report_key": "builtin.channel_mix", "time_range": JUNE})
    assert resp.status_code == 202, resp.text
    assert len(timeouts) == 1
    assert str(timeouts[0]) in ("0", "0.000000", "0.0")


async def test_scope_is_reevaluated_in_job(env: Env) -> None:
    body = {"report_key": "builtin.channel_mix", "time_range": JUNE, "force_job": True,
            "no_cache": True}  # fmt: skip
    granted = (await _run(env, body)).json()["job"]["job_id"]
    await env.run_jobs()
    assert (await env.job(granted))["result"]["rows"] == [["email", 2]]
    # queued while the agent can read the queue, run after the grant is gone
    revoked = (await _run(env, body)).json()["job"]["job_id"]
    _run_sql(env.sync_url, (f"DELETE FROM group_user WHERE user_id = {AGENT}",))
    try:
        await env.run_jobs()
        job = await env.job(revoked)
        assert job["status"] == "succeeded", job
        assert job["result"]["rows"] == []
    finally:
        engine = create_engine(env.sync_url)
        with engine.begin() as conn:
            _grant(conn)
        engine.dispose()


# ---------------------------------------------------------------- exports


async def _export(env: Env, fmt: str) -> tuple[dict[str, Any], bytes]:
    resp = await env.client.post(
        "/api/v1/reports/export",
        json={"report_key": "builtin.time_by_customer", "format": fmt, "title": "Zeit",
              "time_range": JUNE, "column_labels": {"customer": "Kunde"},
              "info_lines": ["Zeitraum: Juni 2024"]},
    )  # fmt: skip
    assert resp.status_code == 202, resp.text
    assert set(resp.json()) == {"job_id", "status"}
    await env.run_jobs()
    job = await env.job(resp.json()["job_id"])
    assert job["status"] == "succeeded", job
    assert job["kind"] == "core.report_export"
    assert len(job["files"]) == 1
    file = job["files"][0]
    dl = await env.client.get(f"/api/v1/jobs/{job['id']}/files/{file['id']}")
    assert dl.status_code == 200
    return job, dl.content


async def test_export_xlsx_job_file(env: Env) -> None:
    job, content = await _export(env, "xlsx")
    assert job["files"][0]["name"] == "zeit.xlsx"
    assert job["files"][0]["content_type"] == XLSX
    assert job["result"] == {"rows": 1}
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Zeit"]
    values = [list(r) for r in ws.iter_rows(values_only=True)]
    assert values[0][0] == "Kunde"
    assert values[1][0] == CUSTOMER and values[1][1] == pytest.approx(3.5)  # foreign 3.0 not
    assert "Info" in wb.sheetnames


async def test_export_csv_job_file(env: Env) -> None:
    job, content = await _export(env, "csv")
    assert job["files"][0]["name"] == "zeit.csv"
    assert job["files"][0]["content_type"] == "text/csv; charset=utf-8"
    assert job["result"] == {"rows": 1}
    assert content.startswith(b"\xef\xbb\xbf")
    lines = content.decode("utf-8-sig").splitlines()
    assert lines[0].startswith("Kunde;")
    assert lines[1].startswith(f"{CUSTOMER};3.5;")


def test_export_slug() -> None:
    assert report_jobs.export_slug("Zeit") == "zeit"
    assert report_jobs.export_slug("Größe & Übersicht 2024") == "groesse-uebersicht-2024"
    assert report_jobs.export_slug("¿?") == "bericht"


# ---------------------------------------------------------------- privacy and limits


async def test_report_job_results_are_private_to_the_creator(env: Env) -> None:
    run = await _run(
        env, {"report_key": "builtin.channel_mix", "time_range": JUNE, "force_job": True}
    )
    run_id = run.json()["job"]["job_id"]
    export = await env.client.post(
        "/api/v1/reports/export",
        json={"report_key": "builtin.time_by_customer", "format": "csv", "time_range": JUNE},
    )
    export_id = export.json()["job_id"]
    await env.run_jobs()
    own = await env.job(run_id)
    assert own["result"]["rows"] == [["email", 2]]
    assert (await env.job(export_id))["result"] == {"rows": 1}
    own_list = (await env.client.get("/api/v1/jobs")).json()
    assert next(j for j in own_list if j["id"] == run_id)["result"] is not None

    env.user = ADMIN
    for job_id in (run_id, export_id):
        seen = await env.job(job_id)  # admins see every job ...
        assert seen["status"] == "succeeded" and seen["created_by"] == AGENT
        assert seen["result"] is None  # ... but not another user's report data
    listed = {j["id"]: j for j in (await env.client.get("/api/v1/jobs?all=true")).json()}
    assert listed[run_id]["result"] is None and listed[export_id]["result"] is None
    file_id = (await env.job(export_id))["files"][0]["id"]
    download = await env.client.get(f"/api/v1/jobs/{export_id}/files/{file_id}")
    assert download.status_code == 404


async def test_report_jobs_are_capped_per_user(env: Env) -> None:
    limit = reports_api.MAX_ACTIVE_REPORT_JOBS_PER_USER
    assert limit == 3
    body = {"report_key": "builtin.channel_mix", "time_range": JUNE, "force_job": True,
            "no_cache": True}  # fmt: skip
    for _ in range(limit - 1):
        assert (await _run(env, body)).status_code == 202
    export_body = {"report_key": "builtin.time_by_customer", "time_range": JUNE}
    assert (await env.client.post("/api/v1/reports/export", json=export_body)).status_code == 202
    over = await _run(env, body)
    assert over.status_code == 429
    assert over.json()["detail"]["message"] == (
        "Es laufen bereits 3 Berichte im Hintergrund. Bitte warte, bis einer fertig ist."
    )
    over_export = await env.client.post("/api/v1/reports/export", json=export_body)
    assert over_export.status_code == 429
    # inline runs are not limited
    inline = await _run(env, {"report_key": "builtin.channel_mix", "time_range": JUNE})
    assert inline.status_code == 200
    # another user has their own allowance
    env.user = ADMIN
    assert (await _run(env, body)).status_code == 202
    env.user = USER
    # finished jobs free the slots
    await env.run_jobs()
    await env.run_jobs()
    assert (await _run(env, body)).status_code == 202


async def test_cancelled_report_job_stops_its_statement_and_ends_cancelled(
    env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Frontend-I1 (backend): a cancel stops a running report job within the poll interval;
    the statement in flight is cut off and its connection invalidated."""
    from sqlalchemy import func, select

    import tiqora.reporting.engine as engine_mod

    monkeypatch.setattr(report_jobs, "CANCEL_POLL_SECONDS", 0.1)
    started = asyncio.Event()
    invalidated: list[bool] = []
    real_invalidate = engine_mod.invalidate_quietly

    async def spy_invalidate(session: AsyncSession) -> None:
        invalidated.append(True)
        await real_invalidate(session)

    async def slow_report(session: AsyncSession, **kwargs: Any) -> Any:
        sleep = func.sleep(30) if env.dialect == "mysql" else func.pg_sleep(30)
        started.set()
        await engine_mod._fetch(session, select(sleep), QueryBudget(per_query_seconds=60))
        raise AssertionError("the statement was not cut off")

    monkeypatch.setattr(engine_mod, "invalidate_quietly", spy_invalidate)
    monkeypatch.setattr(report_jobs, "run_report", slow_report)
    resp = await _run(env, {"report_key": "builtin.channel_mix", "time_range": JUNE,
                            "force_job": True, "no_cache": True})  # fmt: skip
    job_id = resp.json()["job"]["job_id"]
    runner = asyncio.create_task(env.run_jobs())
    await asyncio.wait_for(started.wait(), 15)
    await asyncio.sleep(0.3)  # the statement is now running on the server
    cancelled_at = time.monotonic()
    assert (await env.client.post(f"/api/v1/jobs/{job_id}/cancel")).status_code == 200
    await asyncio.wait_for(runner, 15)
    assert time.monotonic() - cancelled_at < 10  # not the statement's 30 s
    assert (await env.job(job_id))["status"] == "cancelled"
    assert invalidated  # cut off mid-statement: the connection never returns to the pool


async def test_cancel_requested_report_jobs_do_not_count_for_the_cap(env: Env) -> None:
    body = {"report_key": "builtin.channel_mix", "time_range": JUNE, "force_job": True,
            "no_cache": True}  # fmt: skip
    ids = [(await _run(env, body)).json()["job"]["job_id"] for _ in range(3)]
    assert (await _run(env, body)).status_code == 429
    async with env.factory() as s:
        # running and asked to stop: the job ends within a second (see the test above)
        await s.execute(
            text("UPDATE tiqora_job SET status = 'running', cancel_requested = :t WHERE id = :i"),
            {"t": True, "i": ids[0]},
        )
        await s.commit()
    assert (await _run(env, body)).status_code == 202
    assert (await _run(env, body)).status_code == 429


async def test_concurrent_report_jobs_cannot_overshoot_the_cap(
    env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Five background runs at once: the user-row lock lets exactly 3 through."""
    real_count = run_service.count_active_jobs

    async def slow_count(*args: Any, **kwargs: Any) -> int:
        n = await real_count(*args, **kwargs)
        await asyncio.sleep(0.2)  # widen the check-then-insert window
        return n

    monkeypatch.setattr(run_service, "count_active_jobs", slow_count)
    # five planners racing to CREATE the never-committed state row would deadlock in
    # load_state (its documented first-start contract) before reaching the cap
    await env.ensure_state()
    body = {"report_key": "builtin.channel_mix", "time_range": JUNE, "force_job": True,
            "no_cache": True}  # fmt: skip
    responses = await asyncio.gather(*(_run(env, body) for _ in range(5)))
    assert sorted(r.status_code for r in responses) == [202, 202, 202, 429, 429]
    async with env.factory() as s:
        active = (
            await s.execute(
                text(
                    f"SELECT COUNT(*) FROM tiqora_job WHERE {_JOBS}"
                    " AND created_by = :u AND status IN ('queued', 'running')"
                ),
                {"u": AGENT},
            )
        ).scalar_one()
    assert active == 3


async def test_request_size_limits(env: Env) -> None:
    run = {"report_key": "builtin.channel_mix"}
    too_big: list[tuple[str, dict[str, Any]]] = [
        ("parameters", {"parameters": {f"p{i}": 1 for i in range(11)}}),
        ("parameters", {"parameters": {"queues": list(range(1001))}}),
        ("parameters", {"parameters": {"queues": ["x" * 1001]}}),
        ("parameters", {"parameters": {"k" * 41: 1}}),
        ("dimension_choices", {"dimension_choices": {str(i): "type" for i in range(4)}}),
        ("dimension_choices", {"dimension_choices": {"0": "t" * 121}}),
        ("report_key", {"report_key": "k" * 121}),
    ]
    for field, extra in too_big:
        resp = await _run(env, {**run, **extra})
        assert resp.status_code == 422, (field, resp.text)
        assert any(field in err["loc"] for err in resp.json()["detail"]), field
    export_too_big: list[tuple[str, dict[str, Any]]] = [
        ("column_labels", {"column_labels": {f"c{i}": "x" for i in range(21)}}),
        ("column_labels", {"column_labels": {"channel": "x" * 201}}),
        ("value_labels", {"value_labels": {f"c{i}": {} for i in range(21)}}),
        ("value_labels", {"value_labels": {"channel": {str(i): "x" for i in range(501)}}}),
        ("value_labels", {"value_labels": {"channel": {"email": "x" * 201}}}),
        ("info_lines", {"info_lines": ["x"] * 21}),
        ("info_lines", {"info_lines": ["x" * 301]}),
        ("totals_label", {"totals_label": "x" * 41}),
    ]
    for field, extra in export_too_big:
        resp = await env.client.post("/api/v1/reports/export", json={**run, **extra})
        assert resp.status_code == 422, (field, resp.text)
        assert any(field in err["loc"] for err in resp.json()["detail"]), field
    # at the limits everything is accepted
    ok = await env.client.post(
        "/api/v1/reports/export",
        json={**run, "column_labels": {f"c{i}": "x" * 200 for i in range(20)},
              "value_labels": {"channel": {str(i): "x" for i in range(500)}},
              "info_lines": ["x" * 300] * 20},
    )  # fmt: skip
    assert ok.status_code == 202, ok.text


# ---------------------------------------------------------------- access


async def test_access_setting_stats_group(env: Env) -> None:
    queued = await _run(
        env, {"report_key": "builtin.channel_mix", "time_range": JUNE, "force_job": True}
    )
    assert queued.status_code == 202
    await env.set_access("stats_group")
    for method, path, body in (
        ("GET", "/api/v1/reports", None),
        ("GET", "/api/v1/reports/builtin.channel_mix", None),
        ("POST", "/api/v1/reports/run", {"report_key": "builtin.channel_mix"}),
        ("POST", "/api/v1/reports/export", {"report_key": "builtin.channel_mix"}),
        ("GET", "/api/v1/reporting/sources", None),
        ("GET", "/api/v1/reporting/status", None),
    ):
        resp = await env.client.request(method, path, json=body)
        assert resp.status_code == 403, (path, resp.text)
        assert "stats" in resp.json()["detail"]["message"]
    # the job queued before the change checks access again when it runs
    await env.run_jobs()
    job = await env.job(queued.json()["job"]["job_id"])
    assert job["status"] == "failed"
    assert job["error"] == "Für Berichte fehlt die Berechtigung."
    await env.set_access("all_agents")
    assert (await env.client.get("/api/v1/reports")).status_code == 200


async def test_me_can_view_reports_follows_access_setting(env: Env) -> None:
    """``/auth/me`` ``can_view_reports`` uses the rule of ``require_report_access``."""

    async def flag() -> bool:
        resp = await env.client.get("/api/v1/auth/me")
        assert resp.status_code == 200, resp.text
        return resp.json()["can_view_reports"]  # type: ignore[no-any-return]

    assert await flag() is True  # default all_agents
    await env.set_access("stats_group")
    assert await flag() is False  # agent without the stats group
    assert (await env.client.get("/api/v1/reports")).status_code == 403
    env.user = ADMIN
    assert await flag() is True  # admins always may
    env.user = USER
    async with env.factory() as s:
        # ``ro`` on Znuny's seeded "stats" group grants access (Znuny parity)
        stats = (
            await s.execute(text("SELECT id FROM permission_groups WHERE name = 'stats'"))
        ).scalar_one()
        await s.execute(
            text(
                "INSERT INTO group_user (user_id, group_id, permission_key,"
                " create_time, create_by, change_time, change_by)"
                " VALUES (:u, :g, 'ro', :t, 1, :t, 1)"
            ),
            {"u": AGENT, "g": stats, "t": BUILT},
        )
        await s.commit()
    try:
        assert await flag() is True
        assert (await env.client.get("/api/v1/reports")).status_code == 200
    finally:
        async with env.factory() as s:
            await s.execute(
                text("DELETE FROM group_user WHERE user_id = :u AND group_id = :g"),
                {"u": AGENT, "g": stats},
            )
            await s.commit()
        await env.set_access("all_agents")
