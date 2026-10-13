"""CMDB import: mapping targets, file readers and cell conversion (pure, no database)."""

from __future__ import annotations

import io
from datetime import datetime
from zoneinfo import ZoneInfo

import pytest
from openpyxl import Workbook

from tests._cmdb_helpers import ATTRS
from tiqora.cmdb.schemas import MappingRowIn
from tiqora.cmdb.transfer.import_job import (
    CellProblem,
    _Lookups,
    check_mapping,
    convert_cell,
    default_mapping,
    resolve_columns,
)
from tiqora.cmdb.transfer.mapping import (
    ColumnTarget,
    ImportOptions,
    TargetError,
    format_target,
    parse_target,
)
from tiqora.cmdb.transfer.readers import ReadError, read_rows, unguard


@pytest.mark.parametrize(
    "text,expected",
    [
        ("number", ColumnTarget("number")),
        ("attr:Vendor", ColumnTarget("attr", (("Vendor", 1),))),
        (
            "attr:NIC[2].IPAddress[3]",
            ColumnTarget("attr", (("NIC", 2), ("IPAddress", 3))),
        ),
    ],
)
def test_parse_target(text: str, expected: ColumnTarget) -> None:
    assert parse_target(text) == expected and parse_target(format_target(expected)) == expected


@pytest.mark.parametrize("bad", ["attr:", "attr:NIC[0]", "attr:NIC[x]", "foo", "attr:A[101]"])
def test_parse_target_errors(bad: str) -> None:
    with pytest.raises(TargetError):
        parse_target(bad)


def test_format_target_leaves_out_position_one() -> None:
    assert format_target(parse_target("attr:NIC[1].IPAddress[2]")) == "attr:NIC.IPAddress[2]"
    assert parse_target("attr:NIC.IPAddress[2]").dotted == "NIC.IPAddress"


def test_parse_target_depth_limit() -> None:
    with pytest.raises(TargetError):
        parse_target("attr:" + ".".join(f"A{i}" for i in range(6)))


def test_read_rows_csv_with_bom_and_cp1252() -> None:
    data = "﻿number;name\nCMDB4SRV-000001;Müller-Server\n".encode()
    header, rows = read_rows(data, file_format="csv", options=ImportOptions())
    assert header == ["number", "name"] and list(rows) == [["CMDB4SRV-000001", "Müller-Server"]]
    header, rows = read_rows(
        "name\nMüller\n".encode("cp1252"),
        file_format="csv",
        options=ImportOptions(encoding="cp1252"),
    )
    assert header == ["name"] and list(rows) == [["Müller"]]


def test_read_rows_csv_wrong_encoding_is_a_read_error() -> None:
    with pytest.raises(ReadError):
        read_rows("name\nMüller\n".encode("cp1252"), file_format="csv", options=ImportOptions())


def test_read_rows_csv_without_header_row_numbers_the_columns() -> None:
    header, rows = read_rows(
        b"a;b\nc;d\n", file_format="csv", options=ImportOptions(header_row=False)
    )
    assert header == ["1", "2"] and list(rows) == [["a", "b"], ["c", "d"]]


def test_read_rows_csv_other_delimiter() -> None:
    header, rows = read_rows(b"a,b\n1,2\n", file_format="csv", options=ImportOptions(delimiter=","))
    assert header == ["a", "b"] and list(rows) == [["1", "2"]]


def test_read_rows_xlsx_dates_are_datetimes() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["name", "bought", "cpu"])
    ws.append(["cmdb4-x", datetime(2026, 3, 4, 10, 30), 8])
    buf = io.BytesIO()
    wb.save(buf)
    header, rows = read_rows(buf.getvalue(), file_format="xlsx", options=ImportOptions())
    assert header == ["name", "bought", "cpu"]
    assert list(rows) == [["cmdb4-x", datetime(2026, 3, 4, 10, 30), 8]]


def test_read_rows_xlsx_named_sheet_and_garbage() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Eins"
    ws.append(["a"])
    two = wb.create_sheet("Zwei")
    two.append(["b"])
    two.append(["x"])
    buf = io.BytesIO()
    wb.save(buf)
    header, rows = read_rows(
        buf.getvalue(), file_format="xlsx", options=ImportOptions(sheet="Zwei")
    )
    assert header == ["b"] and list(rows) == [["x"]]
    with pytest.raises(ReadError):
        read_rows(buf.getvalue(), file_format="xlsx", options=ImportOptions(sheet="Drei"))
    with pytest.raises(ReadError):
        read_rows(b"not a zip", file_format="xlsx", options=ImportOptions())


@pytest.mark.parametrize(
    "text,expected",
    [
        ('\'=HYPERLINK("http://example.org")', '=HYPERLINK("http://example.org")'),
        ("'+49 123", "+49 123"),
        ("'-5", "-5"),
        ("'@home", "@home"),
        ("'plain", "'plain"),
        ("=1+1", "=1+1"),
        ("'", "'"),
    ],
)
def test_unguard_strips_exactly_the_export_guard(text: str, expected: str) -> None:
    assert unguard(text) == expected


def test_csv_reader_unguards_cells() -> None:
    _, rows = read_rows(
        b'name\n\'=HYPERLINK("http://example.org")\n',
        file_format="csv",
        options=ImportOptions(),
    )
    assert list(rows) == [['=HYPERLINK("http://example.org")']]


# --- mapping resolution ---------------------------------------------------------------------


def _mapping(*pairs: tuple[str | int, str]) -> list[MappingRowIn]:
    return [MappingRowIn(column=c, target=t) for c, t in pairs]


def test_resolve_columns_by_header_index_and_case() -> None:
    cols, problems = resolve_columns(
        ["Number", "name", "x"], _mapping(("number", "number"), (2, "name"), ("x", "ignore"))
    )
    assert problems == []
    assert [(c.index, c.target.kind) for c in cols] == [(0, "number"), (1, "name")]


def test_resolve_columns_reports_unknown_columns_and_targets() -> None:
    _, problems = resolve_columns(["a"], _mapping(("nope", "name"), (5, "name"), ("a", "foo")))
    assert [p.row for p in problems] == [1, 2, 3]
    assert "nope" in problems[0].message


def test_check_mapping_rules() -> None:
    header = ["name", "vendor", "vendor2", "grp", "extra", "ip"]
    cols, problems = resolve_columns(
        header,
        _mapping(
            ("name", "name"),
            ("vendor", "attr:Vendor"),
            ("vendor2", "attr:Vendor"),
            ("grp", "attr:Specs"),
            ("extra", "attr:Nope"),
            ("ip", "attr:NIC[1].IPAddress[5]"),
        ),
    )
    assert problems == []
    found = check_mapping(cols, ATTRS, ImportOptions(identifiers=["name"]))
    messages = " | ".join(p.message for p in found)
    assert "schon der Spalte" in messages  # Vendor twice
    assert "lässt sich nicht importieren" in messages  # group
    assert "„Nope“" in messages
    assert "höchstens 4 Einträge" in messages


def test_check_mapping_identifier_must_be_mapped_and_searchable() -> None:
    cols, _ = resolve_columns(
        ["name", "model"], _mapping(("name", "name"), ("model", "attr:Model"))
    )
    msgs = [p.message for p in check_mapping(cols, ATTRS, ImportOptions(identifiers=["number"]))]
    assert any("Kennung" in m and "keiner Spalte" in m for m in msgs)
    msgs = [
        p.message for p in check_mapping(cols, ATTRS, ImportOptions(identifiers=["attr:Model"]))
    ]
    assert any("nicht durchsuchbar" in m for m in msgs)
    assert check_mapping(cols, ATTRS, ImportOptions(identifiers=["name"])) == []


def test_check_mapping_create_missing_needs_a_name_column() -> None:
    cols, _ = resolve_columns(["number"], _mapping(("number", "number")))
    msgs = [p.message for p in check_mapping(cols, ATTRS, ImportOptions())]
    assert any("Namen" in m for m in msgs)
    assert check_mapping(cols, ATTRS, ImportOptions(create_missing=False)) == []


def test_default_mapping_maps_export_headers_and_ignores_the_rest() -> None:
    rows = default_mapping(["number", "attr:Vendor", "Some Note"])
    assert [(r.column, r.target) for r in rows] == [
        ("number", "number"),
        ("attr:Vendor", "attr:Vendor"),
        ("Some Note", "ignore"),
    ]


# --- cell conversion ------------------------------------------------------------------------


def _attr(key: str):  # type: ignore[no-untyped-def]
    def walk(items):  # type: ignore[no-untyped-def]
        for a in items:
            if a.key == key:
                return a
            found = walk(a.sub)
            if found is not None:
                return found
        return None

    return walk(ATTRS)


def _env() -> _Lookups:
    return _Lookups(tz=ZoneInfo("Europe/Berlin"))


def test_convert_cell_dates_and_datetimes() -> None:
    env = _env()
    date_attr = _attr("Purchase")
    assert convert_cell(date_attr, "04.03.2026", env) == "2026-03-04"
    assert convert_cell(date_attr, "2026-03-04", env) == "2026-03-04"
    assert convert_cell(date_attr, datetime(2026, 3, 4, 10, 0), env) == "2026-03-04"
    assert convert_cell(date_attr, "morgen", env) == "morgen"  # normalize_values reports it
    dt = date_attr.model_copy(update={"type": "datetime"})
    # Berlin is UTC+1 in March, UTC+2 in July
    assert convert_cell(dt, "2026-03-04 10:00", env) == "2026-03-04T09:00:00"
    assert convert_cell(dt, "04.07.2026 10:00:30", env) == "2026-07-04T08:00:30"
    assert convert_cell(dt, datetime(2026, 7, 4, 10, 0), env) == "2026-07-04T08:00:00"


def test_convert_cell_numbers() -> None:
    env = _env()
    assert convert_cell(_attr("CPU"), 8.0, env) == 8
    assert convert_cell(_attr("CPU"), "8", env) == "8"
    assert convert_cell(_attr("RAM"), "16,5", env) == "16.5"
    assert convert_cell(_attr("RAM"), 16.5, env) == "16.5"


def test_convert_cell_lookups_fail_with_the_cell_text() -> None:
    env = _env()
    owner = _attr("Location")  # ci_ref
    env.ci_numbers["CMDB4SRV-000001"] = 7
    assert convert_cell(owner, "CMDB4SRV-000001", env) == 7
    with pytest.raises(CellProblem) as exc:
        convert_cell(owner, "CMDB4SRV-999999", env)
    assert "CMDB4SRV-999999" in str(exc.value)


def test_report_cells_are_formula_guarded() -> None:
    import csv

    from tiqora.cmdb.transfer.import_job import RowPlan, _report

    plan = RowPlan(row_no=2, status="error", number="=1+1", messages=["@SUM(A1)", "-5"])
    text = _report([plan]).decode("utf-8-sig")
    row = next(csv.DictReader(io.StringIO(text), delimiter=";"))
    assert row["ci_number"] == "'=1+1"
    assert row["message"] == "'@SUM(A1); -5"  # only a leading character matters
    assert row["status"] == "error" and row["row"] == "2"
