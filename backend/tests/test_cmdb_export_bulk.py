"""CMDB export (CSV/XLSX) and bulk edit as jobs, both dialects (``cmdb_db``).

Seed band 744500-744999 (names contain ``cmdb4``).
"""

from __future__ import annotations

import csv
import io
import json
from typing import Any

import pytest
from openpyxl import load_workbook

from tests import _cmdb_helpers
from tests._cmdb_helpers import (
    GROUP_RO,
    SERVER_ATTRS,
    SERVICE_ID,
    USER_A,
    CmdbWorld,
    _attr,
    client_for,
    make_ci,
    make_class,
    run_jobs_until_idle,
)
from tiqora.cmdb.transfer import bulk_job as bulk_mod
from tiqora.cmdb.transfer.export_job import ENTRY_CAP
from tiqora.cmdb.transfer.mapping import default_mapping_for_header

pytestmark = pytest.mark.db

cmdb_db = _cmdb_helpers.cmdb_db
_cmdb_rows_reset = _cmdb_helpers._cmdb_rows_reset


def _nic(*pairs: tuple[str, str]) -> dict[str, Any]:
    return {
        "NIC": [{"v": n, "s": {"IPAddress": [{"v": ip}]}} for n, ip in pairs],
    }


async def _three(db: CmdbWorld, **kw: Any) -> tuple[int, list[int]]:
    async with db.factory() as s:
        cid = await make_class(s, **kw)
        ids = [
            await make_ci(
                s,
                cid,
                "cmdb4-e1",
                attributes={"Vendor": [{"v": "Example"}], **_nic(("eth0", "192.0.2.1"))},
            ),
            await make_ci(
                s,
                cid,
                "cmdb4-e2",
                attributes=_nic(("eth0", "192.0.2.2"), ("eth1", "192.0.2.3")),
                deployment="maintenance",
            ),
            await make_ci(s, cid, "cmdb4-e3", attributes={"Vendor": [{"v": "=HYPERLINK(1)"}]}),
        ]
    return cid, ids


async def _export(db: CmdbWorld, cid: int, fmt: str = "csv", **extra: Any) -> int:
    async with client_for(db, USER_A) as c:
        r = await c.post(
            "/api/v1/cmdb/export-jobs", json={"class_id": cid, "file_format": fmt, **extra}
        )
    assert r.status_code == 202, r.text
    return int(r.json()["job_id"])


async def _job(db: CmdbWorld, job_id: int) -> dict[str, Any]:
    await run_jobs_until_idle(db)
    async with client_for(db, USER_A) as c:
        r = await c.get(f"/api/v1/jobs/{job_id}")
    assert r.status_code == 200, r.text
    body: dict[str, Any] = r.json()
    return body


def _file(db: CmdbWorld, job_id: int, name: str) -> bytes:
    rows = db.sql(
        "SELECT data FROM tiqora_job_file WHERE job_id = :j AND name = :n", j=job_id, n=name
    )
    assert len(rows) == 1
    return bytes(rows[0][0])


def _csv_rows(data: bytes) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig")), delimiter=";"))


async def test_export_csv_round_trips_through_import(cmdb_db: CmdbWorld) -> None:
    cid, _ = await _three(cmdb_db)
    job = await _export(cmdb_db, cid)
    res = await _job(cmdb_db, job)
    assert res["status"] == "succeeded" and res["result"]["rows"] == 3, res
    data = _file(cmdb_db, job, "cmdb-cmdb4srv.csv")
    rows = _csv_rows(data)
    assert [r["number"][:8] for r in rows] == ["CMDB4SRV"] * 3
    assert rows[1]["attr:NIC[2].IPAddress"] == "192.0.2.3"
    assert rows[1]["deployment_state"] == "Maintenance"
    # F13: a formula-looking cell is guarded
    assert rows[2]["attr:Vendor"] == "'=HYPERLINK(1)"
    header = next(csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter=";"))
    mapping = default_mapping_for_header(header)
    assert [m.target for m in mapping[:2]] == ["number", "name"]
    assert next(m for m in mapping if m.column == "class").target == "ignore"

    async with client_for(cmdb_db, USER_A) as c:
        r = await c.post(
            "/api/v1/cmdb/import-jobs",
            files={"file": ("cmdb4.csv", data, "text/csv")},
            data={"request": json.dumps({"class_id": cid, "dry_run": False})},
        )
    assert r.status_code == 202, r.text
    imp = await _job(cmdb_db, int(r.json()["job_id"]))
    assert imp["status"] == "succeeded", imp
    assert imp["result"]["unchanged"] == 3
    assert imp["result"]["created"] == 0 and imp["result"]["updated"] == 0


async def test_export_respects_permission_at_run_time(cmdb_db: CmdbWorld) -> None:
    cid, _ = await _three(cmdb_db)
    job = await _export(cmdb_db, cid)
    # rights removed between enqueue and run: the class moves to a group USER_A cannot read
    cmdb_db.sql("UPDATE tiqora_cmdb_class SET group_id = :g WHERE id = :c", g=740012, c=cid)
    res = await _job(cmdb_db, job)
    assert res["status"] == "failed"
    assert "keine Berechtigung" in json.dumps(res, ensure_ascii=False)


async def test_export_xlsx_has_header_and_rows(cmdb_db: CmdbWorld) -> None:
    cid, _ = await _three(cmdb_db)
    job = await _export(cmdb_db, cid, "xlsx", q="cmdb4-e2")
    res = await _job(cmdb_db, job)
    assert res["status"] == "succeeded" and res["result"]["rows"] == 1, res
    wb = load_workbook(io.BytesIO(_file(cmdb_db, job, "cmdb-cmdb4srv.xlsx")), read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))  # type: ignore[union-attr]
    assert list(rows[0][:5]) == ["number", "name", "class", "deployment_state", "incident_state"]
    assert len(rows) == 2 and rows[1][1] == "cmdb4-e2"
    assert "attr:NIC[2].IPAddress" in rows[0]


async def test_export_caps_entry_columns_at_10(cmdb_db: CmdbWorld) -> None:
    attrs = [*SERVER_ATTRS[:1], _attr("Disk", "Disk", "text", count_max=15)]
    async with cmdb_db.factory() as s:
        cid = await make_class(s, attributes=attrs)
        await make_ci(
            s, cid, "cmdb4-many", attributes={"Disk": [{"v": f"d{i}"} for i in range(12)]}
        )
    job = await _export(cmdb_db, cid)
    res = await _job(cmdb_db, job)
    assert res["status"] == "succeeded", res
    assert res["result"]["entries_cut"] is True
    row = _csv_rows(_file(cmdb_db, job, "cmdb-cmdb4srv.csv"))[0]
    disk_cols = [k for k in row if k.startswith("attr:Disk")]
    assert len(disk_cols) == ENTRY_CAP
    assert row["attr:Disk[10]"] == "d9" and "attr:Disk[11]" not in row


async def _bulk(db: CmdbWorld, body: dict[str, Any]) -> int:
    async with client_for(db, USER_A) as c:
        r = await c.post("/api/v1/cmdb/bulk-jobs", json=body)
    assert r.status_code == 202, r.text
    return int(r.json()["job_id"])


async def _states(db: CmdbWorld) -> dict[str, int]:
    async with db.factory() as s:
        return {
            "maintenance": await _cmdb_helpers._dep(s, "maintenance"),
            "incident": await _cmdb_helpers._inc(s, "incident"),
        }


async def test_bulk_sets_state_and_link_and_skips_without_rw(cmdb_db: CmdbWorld) -> None:
    async with cmdb_db.factory() as s:
        rw = await make_class(s)
        ro = await make_class(s, key="cmdb4ro", group_id=GROUP_RO)
        a = await make_ci(s, rw, "cmdb4-b1")
        b = await make_ci(s, rw, "cmdb4-b2")
        c3 = await make_ci(s, ro, "cmdb4-b3")
    st = await _states(cmdb_db)
    job = await _bulk(
        cmdb_db,
        {
            "ci_ids": [a, b, c3],
            "deployment_state_id": st["maintenance"],
            "link": {
                "link_type_key": "DependsOn",
                "target_kind": "service",
                "target_key": str(SERVICE_ID),
            },
        },
    )
    res = await _job(cmdb_db, job)
    assert res["status"] == "succeeded", res
    assert res["result"] == {"updated": 2, "skipped": 1}
    counts = {
        int(r[0]): (int(r[1]), int(r[2]))
        for r in cmdb_db.sql(
            "SELECT id, version_count, deployment_state_id FROM tiqora_cmdb_ci "
            "WHERE id IN (:a, :b, :c)",
            a=a,
            b=b,
            c=c3,
        )
    }
    assert counts[a] == (2, st["maintenance"]) and counts[b] == (2, st["maintenance"])
    assert counts[c3][0] == 1 and counts[c3][1] != st["maintenance"]
    links = cmdb_db.sql(
        "SELECT source_key FROM tiqora_itsm_link WHERE target_kind = 'service' "
        "AND target_key = :s OR source_kind = 'service' AND source_key = :s",
        s=str(SERVICE_ID),
    )
    assert len(links) == 2 and {str(r[0]) for r in links} == {str(a), str(b)}
    report = _csv_rows(_file(cmdb_db, job, "report.csv"))
    assert [r["status"] for r in report] == ["updated", "updated", "skipped"]


async def test_bulk_incident_triggers_one_recompute(
    cmdb_db: CmdbWorld, monkeypatch: pytest.MonkeyPatch
) -> None:
    async with cmdb_db.factory() as s:
        cid = await make_class(s)
        ids = [await make_ci(s, cid, f"cmdb4-r{i}") for i in range(3)]
    st = await _states(cmdb_db)
    calls: list[list[int]] = []
    real = bulk_mod.recompute_from

    async def spy(session: Any, seeds: Any, **kw: Any) -> Any:
        calls.append(sorted(seeds))
        return await real(session, seeds, **kw)

    monkeypatch.setattr(bulk_mod, "recompute_from", spy)
    job = await _bulk(cmdb_db, {"ci_ids": ids, "incident_state_id": st["incident"]})
    res = await _job(cmdb_db, job)
    assert res["status"] == "succeeded" and res["result"]["updated"] == 3, res
    assert calls == [sorted(ids)]


async def test_bulk_over_5000_is_422(cmdb_db: CmdbWorld) -> None:
    async with client_for(cmdb_db, USER_A) as c:
        too_many = await c.post(
            "/api/v1/cmdb/bulk-jobs",
            json={"ci_ids": list(range(1, 5002)), "deployment_state_id": 1},
        )
        nothing = await c.post("/api/v1/cmdb/bulk-jobs", json={"ci_ids": [1]})
        unknown_state = await c.post(
            "/api/v1/cmdb/bulk-jobs", json={"ci_ids": [1], "deployment_state_id": 99999999}
        )
    assert too_many.status_code == 422
    assert nothing.status_code == 422
    assert unknown_state.status_code == 422
