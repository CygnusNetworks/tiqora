"""Readers of the import file: CSV (BOM, encoding, delimiter) and XLSX (read-only).

``read_rows`` returns the header and a lazy iterator of data rows; blank rows are kept (the
job needs the row numbers of the file). XLSX cells keep their type (``datetime``, ``int``,
``float``); CSV cells are strings.
"""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Iterator
from datetime import date, datetime, time
from typing import Any, Final, Literal

from openpyxl import load_workbook

from tiqora.cmdb.transfer.mapping import ImportOptions
from tiqora.reporting.export import FORMULA_PREFIXES

Cell = str | datetime | int | float | None
_BOM: Final = b"\xef\xbb\xbf"
#: guard against zip bombs: a 25 MB upload must not unpack to more than this
_MAX_UNPACKED: Final = 400 * 1024 * 1024


class ReadError(ValueError):
    """The file cannot be read; the message is German and user-facing."""


def unguard(text: str) -> str:
    """Undo the export's formula guard: ``'=SUM(A1)`` becomes ``=SUM(A1)`` again (F13).

    Only an apostrophe in front of ``= + - @ TAB CR`` is removed, nothing else.
    """
    if len(text) > 1 and text[0] == "'" and text[1] in FORMULA_PREFIXES:
        return text[1:]
    return text


def _csv_rows(data: bytes, options: ImportOptions) -> tuple[list[str], Iterator[list[Cell]]]:
    if data.startswith(_BOM):
        data = data[len(_BOM) :]
    try:
        text = data.decode(options.encoding)
    except UnicodeDecodeError as exc:
        raise ReadError(
            "Die Datei passt nicht zur gewählten Zeichenkodierung. "
            "Bitte wählen Sie die andere Kodierung."
        ) from exc
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=options.delimiter)

    def rows() -> Iterator[list[Cell]]:
        try:
            for row in reader:
                yield [unguard(c) for c in row]
        except csv.Error as exc:
            raise ReadError(f"Die CSV-Datei ist fehlerhaft: {exc}") from exc

    return _split_header(rows(), options)


def _xlsx_rows(data: bytes, options: ImportOptions) -> tuple[list[str], Iterator[list[Cell]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > _MAX_UNPACKED:
                raise ReadError("Die Excel-Datei ist entpackt zu groß.")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except ReadError:
        raise
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for broken files
        raise ReadError("Die Datei ist keine gültige Excel-Datei (.xlsx).") from exc
    if options.sheet:
        if options.sheet not in wb.sheetnames:
            wb.close()
            raise ReadError(f"Das Tabellenblatt „{options.sheet}“ gibt es nicht.")
        ws = wb[options.sheet]
    else:
        ws = wb.worksheets[0]

    def rows() -> Iterator[list[Cell]]:
        try:
            for raw in ws.iter_rows(values_only=True):
                yield [_xlsx_cell(c) for c in raw]
        finally:
            wb.close()

    return _split_header(rows(), options)


def _xlsx_cell(value: Any) -> Cell:
    if value is None:
        return None
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, int | float):
        return value
    return unguard(str(value))


def _split_header(
    stream: Iterator[list[Cell]], options: ImportOptions
) -> tuple[list[str], Iterator[list[Cell]]]:
    first = next(stream, None)
    if first is None:
        return [], iter(())
    if options.header_row:
        return [("" if c is None else str(c)).strip() for c in first], stream
    header = [str(i) for i in range(1, len(first) + 1)]

    def chained() -> Iterator[list[Cell]]:
        yield first
        yield from stream

    return header, chained()


def read_rows(
    data: bytes, *, file_format: Literal["csv", "xlsx"], options: ImportOptions
) -> tuple[list[str], Iterator[list[Cell]]]:
    """Header (or ``["1", "2", ...]`` without a header row) and the data rows."""
    if file_format == "csv":
        return _csv_rows(data, options)
    return _xlsx_rows(data, options)


__all__ = ["Cell", "ReadError", "read_rows", "unguard"]
