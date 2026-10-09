"""XLSX and CSV writers for report results (pure, no database)."""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from itertools import chain, islice
from typing import Final

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from tiqora.reporting.schemas import Cell, ColumnType, ReportResultOut

OTHER_KEY: Final = "__other__"
OTHER_LABEL: Final = "Weitere"
_WIDTH_SCAN_ROWS: Final = 500
_ILLEGAL_XML: Final = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_SHEET_FORBIDDEN: Final = re.compile(r"[\[\]:*?/\\]")
_FORMULA_PREFIXES: Final = ("=", "+", "-", "@", "\t", "\r")
_NUMBER_FORMATS: Final[dict[str, str]] = {
    "count": "0",
    "id": "0",
    "hour": "0",
    "weekday": "0",
    "number": "0.00",
    "hours": "0.00",
    "days": "0.00",
    "units": "0.00",
    "minutes": "0",
    "percent": "0.0%",
    "date": "yyyy-mm-dd",
}
_HEADER_FILL: Final = PatternFill("solid", fgColor="EDF0F7")


@dataclass(frozen=True, slots=True)
class XlsxColumn:
    header: str
    type: ColumnType
    width: float | None = None


@dataclass(frozen=True, slots=True)
class XlsxSheet:
    #: sanitised: []:*?/\ removed, at most 31 chars, unique per workbook
    name: str
    columns: Sequence[XlsxColumn]
    rows: Iterable[Sequence[Cell | date]]
    totals: Sequence[Cell] | None = None
    totals_label: str = "Summe"
    #: written to an "Info" sheet when non-empty
    info_lines: Sequence[str] = ()


def _sheet_name(raw: str, used: set[str]) -> str:
    base = _SHEET_FORBIDDEN.sub("", raw).strip().strip("'") or "Sheet"
    name = base[:31]
    n = 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def _text(value: str) -> str:
    return _ILLEGAL_XML.sub("", value)


def _to_date(value: object) -> date | datetime | None:
    if isinstance(value, datetime | date):
        return value
    if isinstance(value, str):
        try:
            return (
                datetime.fromisoformat(value)
                if len(value) > 10
                else datetime.combine(date.fromisoformat(value), datetime.min.time())
            )
        except ValueError:
            return None
    return None


def _make_cell(
    ws: object, value: object, col_type: ColumnType, *, bold: bool = False
) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws)
    if value is None:
        cell.value = None
    elif col_type == "date" and (parsed := _to_date(value)) is not None:
        cell.value = parsed
        cell.number_format = _NUMBER_FORMATS["date"]
    elif isinstance(value, bool):
        cell.value = value
    elif isinstance(value, int | float):
        cell.value = value
        cell.number_format = _NUMBER_FORMATS.get(col_type, "General")
    else:
        cell.value = _text(str(value))
        # never let a text starting with "=" become a formula
        cell.data_type = "s"
    if bold:
        cell.font = Font(bold=True)
    return cell


def _display_len(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, datetime | date):
        return 10
    return len(str(value))


def write_xlsx_workbook(sheets: Sequence[XlsxSheet]) -> bytes:
    wb = Workbook(write_only=True)
    used: set[str] = set()
    info: list[str] = []
    for sheet in sheets:
        ws = wb.create_sheet(_sheet_name(sheet.name, used))
        rows_iter = iter(sheet.rows)
        head = list(islice(rows_iter, _WIDTH_SCAN_ROWS))
        scan = list(head)
        if sheet.totals is not None:
            scan.append(sheet.totals)
        for idx, col in enumerate(sheet.columns):
            if col.width is not None:
                width = col.width
            else:
                longest = max(
                    [len(col.header), *(_display_len(r[idx]) for r in scan if idx < len(r))]
                )
                width = min(60.0, max(10.0, longest + 2.0))
            ws.column_dimensions[get_column_letter(idx + 1)].width = width
        ws.freeze_panes = "A2"
        header: list[WriteOnlyCell] = []
        for col in sheet.columns:
            cell = _make_cell(ws, col.header, "string", bold=True)
            cell.fill = _HEADER_FILL
            header.append(cell)
        ws.append(header)
        for row in chain(head, rows_iter):
            ws.append([_make_cell(ws, v, c.type) for v, c in zip(row, sheet.columns, strict=False)])
        if sheet.totals is not None:
            totals = list(sheet.totals)
            for i, v in enumerate(totals):
                if v is None:
                    totals[i] = sheet.totals_label
                    break
            ws.append(
                [
                    _make_cell(ws, v, c.type, bold=True)
                    for v, c in zip(totals, sheet.columns, strict=False)
                ]
            )
        info.extend(sheet.info_lines)
    if info:
        ws = wb.create_sheet(_sheet_name("Info", used))
        ws.column_dimensions["A"].width = min(60.0, max(10.0, max(len(x) for x in info) + 2.0))
        for line in info:
            ws.append([_make_cell(ws, line, "string")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _dimension_label(
    result: ReportResultOut,
    key: str,
    value: Cell,
    value_labels: Mapping[str, Mapping[str, str]],
) -> Cell:
    if value is None:
        return None
    text = str(value)
    own = value_labels.get(key, {})
    if text in own:
        return own[text]
    if text == OTHER_KEY:
        return OTHER_LABEL
    return result.labels.get(key, {}).get(text, value)


def _labelled_rows(
    result: ReportResultOut, value_labels: Mapping[str, Mapping[str, str]]
) -> list[list[Cell]]:
    out: list[list[Cell]] = []
    for row in result.rows:
        cells: list[Cell] = []
        for col, v in zip(result.columns, row, strict=False):
            # date values stay ISO (real date cells) unless a label maps them
            keep = col.role != "dimension" or (
                col.type == "date"
                and str(v)
                not in {**result.labels.get(col.key, {}), **value_labels.get(col.key, {})}
            )
            cells.append(v if keep else _dimension_label(result, col.key, v, value_labels))
        out.append(cells)
    return out


def report_sheet(
    result: ReportResultOut,
    *,
    title: str,
    column_labels: Mapping[str, str],
    value_labels: Mapping[str, Mapping[str, str]],
    totals_label: str,
    info_lines: Sequence[str],
) -> XlsxSheet:
    columns = [XlsxColumn(column_labels.get(c.key, c.key), c.type) for c in result.columns]
    rows = _labelled_rows(result, value_labels)
    totals = result.totals.grand if result.totals is not None else None
    return XlsxSheet(
        name=title,
        columns=columns,
        rows=rows,
        totals=totals,
        totals_label=totals_label,
        info_lines=info_lines,
    )


def _csv_cell(value: Cell) -> str:
    if value is None:
        return ""
    text = str(value)
    if isinstance(value, str) and text.startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def write_report_csv(
    result: ReportResultOut,
    *,
    column_labels: Mapping[str, str],
    value_labels: Mapping[str, Mapping[str, str]],
    totals_label: str,
) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    writer.writerow([_csv_cell(column_labels.get(c.key, c.key)) for c in result.columns])
    for row in _labelled_rows(result, value_labels):
        writer.writerow([_csv_cell(v) for v in row])
    if result.totals is not None and result.totals.grand is not None:
        grand = list(result.totals.grand)
        for i, v in enumerate(grand):
            if v is None:
                grand[i] = totals_label
                break
        writer.writerow([_csv_cell(v) for v in grand])
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")
